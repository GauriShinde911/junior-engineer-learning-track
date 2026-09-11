"""
independent/excel_report_pipeline/pipeline.py

An end‑to‑end pipeline that takes the raw ``sales_messy.csv`` source,
validates it against the ``SALES_SCHEMA`` (from 3.5), cleans it (3.2),
produces the monthly summary (3.3) and the full Excel report (3.4).
If validation fails, the pipeline writes an *Error Summary* sheet into the
Excel workbook listing each validation error and the rows it affects.

Running the script directly will generate ``sample_data/sales_report_with_errors.xlsx``.
"""

from pathlib import Path
import sys
import pandas as pd
import openpyxl

# Add the exercise package paths so we can import the utility modules.
SKILL_ROOT = Path(__file__).resolve().parents[2]
EXERCISES_ROOT = SKILL_ROOT / "exercises"
sys.path.insert(0, str(EXERCISES_ROOT / "3.2-cleaning"))
sys.path.insert(0, str(EXERCISES_ROOT / "3.3-transformations"))
sys.path.insert(0, str(EXERCISES_ROOT / "3.4-excel-workbooks"))
sys.path.insert(0, str(EXERCISES_ROOT / "3.5-validation"))

from clean_messy_dataset import clean, RAW_FILE
from excel_validator import validate, SALES_SCHEMA, ValidationResult
from monthly_summary import build_monthly_summary
from build_excel_report import build_workbook

OUTPUT_FILE = SKILL_ROOT / "sample_data" / "sales_report_with_errors.xlsx"


def add_error_summary_sheet(wb: openpyxl.Workbook, validation: ValidationResult) -> None:
    """Create a human‑readable *Error Summary* sheet based on ValidationResult."""
    ws = wb.create_sheet("Error Summary")
    ws.append(["Column", "Rule", "Message", "Affected Rows (index)"])
    # Apply header styling – reuse same style as other sheets for consistency.
    for cell in ws[1]:
        cell.font = openpyxl.styles.Font(bold=True, color="FFFFFF")
        cell.fill = openpyxl.styles.PatternFill(start_color="C00000", end_color="C00000", fill_type="solid")
        cell.alignment = openpyxl.styles.Alignment(horizontal="center", vertical="center", wrap_text=True)

    for err in validation.errors:
        rows = ", ".join(str(r) for r in err.affected_rows) if err.affected_rows else "-"
        ws.append([err.column, err.rule, err.message, rows])

    # Auto‑fit column widths (approximate).
    for col in ws.columns:
        max_len = max(len(str(cell.value)) if cell.value else 0 for cell in col)
        ws.column_dimensions[openpyxl.utils.get_column_letter(col[0].column)].width = max_len + 2

    ws.freeze_panes = "A2"


def run_pipeline() -> None:
    # 1️⃣ Load raw data
    raw_df = pd.read_csv(RAW_FILE)

    # 2️⃣ Validate raw data – we keep the errors for reporting.
    validation = validate(raw_df, SALES_SCHEMA)

    # 3️⃣ Clean the data (even if validation failed – we still try to produce a report).
    cleaned_df = clean(RAW_FILE)

    # 4️⃣ Build the workbook from the cleaned data.
    products_df = pd.read_csv(SKILL_ROOT / "sample_data" / "products.csv")
    wb = build_workbook(cleaned_df, products_df)

    # 5️⃣ If there were validation errors, add an extra sheet.
    if not validation.passed:
        add_error_summary_sheet(wb, validation)
        print(f"⚠ Validation produced {len(validation.errors)} error(s) – added 'Error Summary' sheet.")
    else:
        print("✅ Validation passed – no error sheet added.")

    # 6️⃣ Persist the workbook.
    wb.save(OUTPUT_FILE)
    print(f"🗒 Workbook written → {OUTPUT_FILE}")


if __name__ == "__main__":
    run_pipeline()
