"""Unit tests for independent excel_report_pipeline."""

import sys
from pathlib import Path
import openpyxl

TEST_DIR = Path(__file__).resolve().parent
MODULE_DIR = TEST_DIR.parent
PIPELINE_DIR = MODULE_DIR / "independent" / "excel_report_pipeline"

sys.path.insert(0, str(MODULE_DIR))
sys.path.insert(0, str(PIPELINE_DIR))

from pipeline import run_pipeline, OUTPUT_FILE


def test_pipeline_execution():
    run_pipeline()
    assert OUTPUT_FILE.exists()
    wb = openpyxl.load_workbook(OUTPUT_FILE)
    assert "Raw Sales" in wb.sheetnames
    assert "Monthly" in wb.sheetnames
    assert "Category Pivot" in wb.sheetnames
    assert "Error Summary" in wb.sheetnames
