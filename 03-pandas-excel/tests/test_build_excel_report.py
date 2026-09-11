"""Unit tests for 3.4 Excel Workbooks."""

import sys
from pathlib import Path
import openpyxl

TEST_DIR = Path(__file__).resolve().parent
MODULE_DIR = TEST_DIR.parent
SUBSECTION_DIR = MODULE_DIR / "exercises" / "3.4-excel-workbooks"
TRANS_DIR = MODULE_DIR / "exercises" / "3.3-transformations"
CLEANING_DIR = MODULE_DIR / "exercises" / "3.2-cleaning"

sys.path.insert(0, str(MODULE_DIR))
sys.path.insert(0, str(SUBSECTION_DIR))
sys.path.insert(0, str(TRANS_DIR))
sys.path.insert(0, str(CLEANING_DIR))

from build_excel_report import build_workbook
from clean_messy_dataset import clean, RAW_FILE
from cross_table_report import load_products


def test_build_workbook():
    sales_df = clean(RAW_FILE)
    products_df = load_products()
    wb = build_workbook(sales_df, products_df)
    assert isinstance(wb, openpyxl.Workbook)
    assert "Raw Sales" in wb.sheetnames
    assert "Monthly" in wb.sheetnames
    assert "Category Pivot" in wb.sheetnames
